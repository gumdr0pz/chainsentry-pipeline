#!/usr/bin/env python3
"""Generate ChainSentry Product Documentation Gaps - Management Review (funding-framed) DOCX.
Data source: machine-extracted gap matrix (cs-gap-extract.json) from MVP_SPECIFICATION.md.
House rule: digests built programmatically from source artifacts."""
import json
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DATA = json.load(open("/home/mtadmin/.hermes/cache/cs-gap-extract.json"))
OUT = "/home/mtadmin/.hermes/intel/chainsentry-pipeline/reports/ChainSentry-Product-Gaps-Management-Review-20261028.docx"

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)

def set_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def para(text, bold=False, size=None, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs: r.font.color.rgb = RGBColor.from_string("1F3864")
    return h

def md_runs(p, text):
    parts = text.split("**")
    for i, chunk in enumerate(parts):
        if not chunk: continue
        r = p.add_run(chunk); r.bold = (i % 2 == 1)

# ---- Title block
para("chain:SENTRY", bold=True, size=20, color="1F3864", space_after=2)
para("Product Documentation Gaps — Management Review & Funding Baseline", bold=True, size=14, space_after=2)
para("Compiled 28 October 2026 (MYT) · v1.0 · CONFIDENTIAL — internal management review", size=9, color="666666", space_after=10)

# ---- Funding abstract (announcement-ready)
heading("1. Announcement-Ready Abstract (for the funding request)", 1)
abstract = ("chain:SENTRY, Aras Integrasi's sovereign blockchain-investigation platform, has completed its "
 "first law-enforcement briefing cycle: PDRM CCID/JSJK Forensik received the full product briefing on 27 October 2026, "
 "an adoption decision is pending, and a government evaluation channel (CyberSecurity Malaysia) opens with a confirmed "
 "11 November leadership meeting. The platform's verified core is strong - 26/26 page routes and 32/32 authenticated "
 "API paths returned live data under independent verification, backed by a 583,574-tag attribution corpus across 78 chains "
 "and sanctions screening at sub-20-millisecond latency. Thirteen documented work packages separate this verified core "
 "from pilot delivery: 4 critical, 3 must-have, 4 should-have, 2 future. The immediate funding need is delivery capacity - "
 "the adopted 7-FTE operating model (CC-01 envelope MYR 100-150k/month) remains unfilled and a 2-engineer span of "
 "responsibility now gates three concurrent government adoption tracks. Positioning is not verification: no external "
 "capability claim is made beyond the evidence base cited herein.")
p = doc.add_paragraph(); md_runs(p, abstract)
p.paragraph_format.space_after = Pt(10)

# ---- Readiness verdict
heading("2. Readiness Verdict (spec-evidenced)", 1)
p = doc.add_paragraph(); md_runs(p, "**Not ready for production/GA — pilot scope only, after Phase 0.** [T2, spec §5.2–5.4]")
for b in [
    "Deployment parity: 22-commit / 32-day development-to-deployment lag (TEC-012) — release risk accumulates daily.",
    "External access: service unreachable externally; named pilot accounts not issued; session URL requires correction (BLK-002/003).",
    "Credential closure: four provider keys previously exposed; rotation outstanding (BLK-001). C1 recorded closed on operator confirmation 8 Sep; evidence item (fingerprint record) remains open in the C2 pack.",
    "Live-vs-demo boundary: a 31-tool registry remains fixture-backed while the dossier registry is live (DEF-009) — no pilot result may be fixture data.",
    "Governance context: reduced-screening interim posture per DAF Option B ruling 26 Oct (DEC-20261026-001); consolidated Phase-0 (M2–M7) status report due 2 Nov 2026 (ACT-20261026-002)."]:
    doc.add_paragraph(b, style="List Bullet")

# ---- Verified strengths (funding story)
heading("3. Verified Strengths (the asset being funded)", 1)
for b in [
    "Verification sweep: 26/26 page routes correct; 32/32 authenticated GET paths returned live data; 44 API handlers in the recorded baseline.",
    "Risk discrimination confirmed: exchange 16/Minimal; OFAC-listed 90/Critical; mixer router 65/High.",
    "Measured performance (warm): screening/attribution <20 ms; trace-to-cash-out ~32 s; composite score 1.4–1.5 s.",
    "Attribution corpus: 583,574 tags / 524,656 addresses / 78 chains; sanctions registry 4,633 addresses; 1,790 checksum-verified dark-web contributions labelled 'observed' — never criminal attribution.",
    "Commercial layer registered: editions-only SKU catalog (Professional RM 96k/yr · Enterprise RM 348k/yr; reference deals RM 698k and RM 300k Y1)."]:
    doc.add_paragraph(b, style="List Bullet")

# ---- Gap register table (from machine extraction)
heading("4. Gap Register — 13 Documented Work Packages", 1)
para("Extracted programmatically from the MVP Specification v3.0 priority gap matrix (§5.4.1). Tier labels are the spec's own; remediation phrasing condensed from the spec's resolution column.", size=9, color="666666", space_after=6)
tbl = doc.add_table(rows=1, cols=4); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, h in enumerate(["Tier", "Work package", "Ref IDs", "What closing it delivers"]):
    r = hdr[i].paragraphs[0].add_run(h); r.bold = True; set_bg(hdr[i], "1F3864")
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
tier_color = {"Critical": "C00000", "Must": "BF8F00", "Should": "2E74B5", "Future": "7F7F7F"}
for row in DATA["gap_matrix"]:
    c = tbl.add_row().cells
    c[0].text = row["priority"]; c[1].text = row["area"]
    c[2].text = ", ".join(row["gap_ids"]) if row["gap_ids"] else "(contract parity)"
    res = row["resolution"]
    if len(res) > 220: res = res[:217].rstrip() + "..."
    c[3].text = res
    set_bg(c[0], tier_color.get(row["priority"], "FFFFFF"))
    r0 = c[0].paragraphs[0].runs[0]; r0.font.color.rgb = RGBColor.from_string("FFFFFF"); r0.bold = True
widths = [Cm(2.0), Cm(4.6), Cm(3.0), Cm(7.4)]
for row in tbl.rows:
    for i, w in enumerate(widths): row.cells[i].width = w

# ---- Funding structure
heading("5. Funding Structure (three tranches)", 1)
p = doc.add_paragraph(); md_runs(p, "**Tranche 1 — Pilot-critical (immediate).** The 4 Critical + 3 Must packages (§4, rows 1–7): release reconciliation, external TLS + named accounts, credential rotation, truth boundary, API contract regeneration, source-role renewal, and blocking quality gates. This tranche is the gate to the Phase-1 exit criterion: a partner-agency investigator completing a case intake-to-report on live intelligence. It also unblocks the stalled first government evaluation (testing-access action open since 15 Oct) and directly serves the pending PDRM adoption decision.")
p = doc.add_paragraph(); md_runs(p, "**Tranche 2 — Productisation.** The 4 Should packages: migration runner + tested backup/restore, second-tenant isolation and quotas, batch screening + signed PDF exports + audit UI, email alerting and observability. These convert a verified pilot into a defensible multi-agency service — the difference between one agency using chain:SENTRY and several.")
p = doc.add_paragraph(); md_runs(p, "**Tranche 3 — Capability expansion (roadmap).** The 2 Future packages: non-EVM tracing beginning with Bitcoin (the PDRM-relevant USDT/Tron line rides the same work), and cryptographically signed, fully-provenanced evidence dossiers — the court-grade assurance tier.")
p = doc.add_paragraph(); md_runs(p, "**Capacity basis:** funding is requested against the adopted 7-FTE operating model (CC-01 envelope MYR 100–150k/month, DEC-20260908-001). The binding constraint is not technology — all 13 packages are documented engineering work against an already-verified core — but hands: a 2-FTE technical unit currently spans three products. The chain:SENTRY hiring exception remains dormant by ruling; this review is the vehicle to activate it.")

# ---- Risk if unfunded
heading("6. Risk if Unfunded", 1)
for b in [
    "Adoption gates open before readiness: PDRM decision (imminent), CSM evaluation (11 Nov anchor), panel launch (composition-gated) — all three can arrive while the product remains pilot-blocked and externally unreachable.",
    "Single point of failure compounds: one engineer carrying deployment parity, security remediation (7 open findings from the 23 Oct full-zone scan) and two sibling products.",
    "Hiring window already lapsed: the Oct 13–20 start window expired unused; every week of delay compounds the post-expiry disposition debt and pushes pilot readiness into Q1 2027.",
    "Competitive window: PDRM is actively shopping for blockchain-detection tooling (incumbent China-sourced; prior exposure to TRM Labs, Chainalysis, Crystal, Uppsala). A funded sovereign-local alternative meets a live procurement gap; an unfunded one watches it close."]:
    doc.add_paragraph(b, style="List Bullet")

# ---- Methodology
heading("7. Methodology & Sources", 1)
for b in [
    "Gap register machine-extracted from products/chainsentry/MVP_SPECIFICATION.md (MVP Product Specification v3.0, registered DOC-20260814-001) — 13 rows, 13/13 baseline figures verified at extraction.",
    "Commercial figures from SKU Catalog v3.0 (DOC-20260924-010). Governance state from DEC-20261026-001, ACT-20261026-002, AIP Gate Tracker (28 Oct check). Adoption posture from CONV-20261028-004 (CCID/JSJK) and EVT-20261027-001 (CSM 11 Nov).",
    "Claim discipline: positioning ≠ verification; no capability claim beyond the cited evidence base; Solana/coverage phrasing corrected to 'attribution + sanctions (tracing: EVM only)' per spec §6 boundaries."]:
    doc.add_paragraph(b, style="List Bullet")

doc.save(OUT)
print("saved", OUT)
